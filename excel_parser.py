"""
excel_parser.py
Windows-safe Excel parser (uses BytesIO — no temp files, no file locking).

Supports:
  Junior  – Grades 6-9  (class sheets A-H, fixed column layout)
  Senior  – Grades 10-11 O/L (dual student per row, wide sheet)
"""
import io, re
from openpyxl import load_workbook

CLASS_SHEETS = set("ABCDEFGH")

JUNIOR_COLS = {
    2: "Sinhala Language", 3: "Tamil Language",
    4: "Buddhism", 5: "Shaivism", 6: "Catholic Doctrine", 7: "Christianity",
    8: "English", 9: "Mathematics", 10: "Science", 11: "History", 12: "Geography",
    13: "Life Skills", 14: "Music (Western)", 15: "Music (Oriental)",
    16: "Art", 17: "Dance", 18: "Drama", 19: "ICT",
    20: "Practical & Technical Skills", 21: "Health & Physical Education",
    22: "Second Language (Sinhala)", 23: "Second Language (Tamil)",
}

OL_LEFT = {
    4: "Sinhala Language", 5: "Tamil Language",
    6: "Catholic Doctrine", 7: "Buddhism",
    10: "English", 11: "Mathematics", 12: "Science", 13: "History",
    14: "Business & Accounting", 15: "Civic Education",
    16: "Second Language (Tamil)", 22: "Music (Oriental)",
    23: "Music (Western)", 24: "Art", 25: "Dance",
    26: "Literature (English)", 27: "Literature (Sinhala)",
    28: "Drama & Performing Arts", 30: "ICT",
    31: "Agriculture & Food Technology", 32: "Health & Physical Education",
    33: "Media Studies",
}


def _to_bytes(source, filename="") -> tuple[bytes, str]:
    """Normalise source to raw bytes + filename string."""
    if isinstance(source, (str, bytes, os.PathLike)):
        import os as _os
        fname = str(source)
        with open(source, "rb") as f:
            return f.read(), fname
    # file-like (BytesIO / UploadedFile)
    source.seek(0)
    return source.read(), filename


import os
def _make_wb(raw: bytes):
    return load_workbook(io.BytesIO(raw), read_only=True, data_only=True)


def _safe_mark(val):
    if val is None: return None
    if isinstance(val, (int, float)):
        f = float(val)
        return round(f, 1) if 0 <= f <= 100 else None
    s = str(val).strip().lower()
    if s in ("ab", "-", "+", ""): return None
    try:
        f = float(s)
        return round(f, 1) if 0 <= f <= 100 else None
    except ValueError:
        return None


def _is_stu(row, off=0):
    try:
        no   = row[off]
        name = row[off + 1]
        return (isinstance(no, (int, float)) and int(no) == no and no > 0
                and isinstance(name, str) and len(name.strip()) > 1)
    except Exception:
        return False


def _parse_marks(row, col_map):
    out = {}
    for ci, subj in col_map.items():
        if ci < len(row):
            v = _safe_mark(row[ci])
            if v is not None:
                out[subj] = v
    return out


def _grade_from_filename(fname: str):
    m = re.search(r'grade[_\s]*(\d+)', fname, re.IGNORECASE)
    if m: return int(m.group(1))
    m = re.search(r'[^0-9](\d{1,2})[A-Za-z]', fname)
    if m:
        g = int(m.group(1))
        if 6 <= g <= 13: return g
    return None


def _year_term(rows):
    year, term = None, 1
    for row in rows[:8]:
        for val in row:
            if isinstance(val, (int, float)) and 2020 <= val <= 2035:
                year = int(val)
            if isinstance(val, str):
                tl = val.lower()
                if 'first'  in tl or 'පළමු' in tl: term = 1
                elif 'second' in tl or 'දෙවන' in tl: term = 2
                elif 'third'  in tl or 'තෙවන' in tl: term = 3
    return year, term


def _find_right_offset(student_rows):
    for row in student_rows[:10]:
        for start in range(100, min(250, len(row))):
            try:
                if (isinstance(row[start], (int, float)) and 1 <= row[start] <= 60
                        and start + 1 < len(row)
                        and isinstance(row[start + 1], str)
                        and len(row[start + 1].strip()) > 3):
                    return start
            except Exception:
                pass
    return 156


# ── Junior (grades 6-9) ───────────────────────────────────────────────────────
def _parse_junior(raw: bytes, hint_grade=None):
    wb  = _make_wb(raw)
    out = []
    for sname in wb.sheetnames:
        if sname.strip().upper() not in CLASS_SHEETS:
            continue
        ws   = wb[sname]
        rows = list(ws.iter_rows(values_only=True))

        grade, cls = hint_grade, sname.strip().upper()
        if len(rows) > 2:
            hr = rows[2]
            if len(hr) > 6 and isinstance(hr[6], (int, float)) and 6 <= hr[6] <= 13:
                grade = int(hr[6])
            if len(hr) > 7 and isinstance(hr[7], str) and hr[7].strip().upper() in CLASS_SHEETS:
                cls = hr[7].strip().upper()

        year, term = _year_term(rows)
        students   = []
        for row in rows:
            if _is_stu(row):
                marks = _parse_marks(row, JUNIOR_COLS)
                if marks:
                    students.append({"seq_no": int(row[0]), "name": str(row[1]).strip(), "marks": marks})
        if students:
            out.append({"class_section": cls, "grade": grade, "term": term, "year": year, "students": students})
    wb.close()
    return out


# ── Senior O/L (grades 10-11) ─────────────────────────────────────────────────
def _parse_senior(raw: bytes, hint_grade=None):
    wb  = _make_wb(raw)
    out = []
    for sname in wb.sheetnames:
        if sname.strip().upper() not in CLASS_SHEETS:
            continue
        ws   = wb[sname]
        rows = list(ws.iter_rows(values_only=True))

        grade, cls = hint_grade, sname.strip().upper()
        for row in rows[:8]:
            for ci, val in enumerate(row):
                if isinstance(val, (int, float)) and 10 <= val <= 13:
                    nxt = row[ci + 1] if ci + 1 < len(row) else None
                    if isinstance(nxt, str) and nxt.strip().upper() in CLASS_SHEETS:
                        grade = int(val); cls = nxt.strip().upper(); break

        year, term = _year_term(rows)
        stu_rows   = [r for r in rows if _is_stu(r)]
        offset     = _find_right_offset(stu_rows) if stu_rows else 156
        OL_RIGHT   = {ci + offset: subj for ci, subj in OL_LEFT.items()}

        students, seen = [], set()

        def _add(no, name, marks):
            key = (int(no), name[:12])
            if key not in seen and marks:
                seen.add(key)
                students.append({"seq_no": int(no), "name": name, "marks": marks})

        for row in rows:
            if _is_stu(row):
                _add(row[0], str(row[1]).strip(), _parse_marks(row, OL_LEFT))
            if len(row) > offset + 1:
                rs = row[offset:]
                if _is_stu(rs):
                    _add(rs[0], str(rs[1]).strip(), _parse_marks(row, OL_RIGHT))

        if students:
            out.append({"class_section": cls, "grade": grade, "term": term, "year": year, "students": students})
    wb.close()
    return out


# ── Public API ────────────────────────────────────────────────────────────────
def detect_and_parse(source, hint_grade=None, filename=""):
    """
    source: file path (str/Path) OR file-like object (BytesIO / Streamlit UploadedFile).
    No temp files created — Windows-safe.
    """
    raw, fname = _to_bytes(source, filename)

    if hint_grade is None:
        hint_grade = _grade_from_filename(fname)

    # Detect format by checking row width
    wb = _make_wb(raw)
    senior = False
    for sname in wb.sheetnames:
        if sname.strip().upper() not in CLASS_SHEETS:
            continue
        ws = wb[sname]
        for row in ws.iter_rows(values_only=True):
            if _is_stu(row) and len(row) > 100:
                senior = True
                break
        break
    wb.close()

    if senior or (hint_grade and hint_grade >= 10):
        return "senior", _parse_senior(raw, hint_grade)
    return "junior", _parse_junior(raw, hint_grade)
